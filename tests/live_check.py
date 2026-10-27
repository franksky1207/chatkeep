"""Opt-in public-page verification. Prints aggregate evidence only, stores nothing.

This checks extraction against the actual decoded source and Word round-tripping.
It does NOT substitute for GPT Export reference or visible browser DOM comparison.
"""
from collections import Counter
from io import BytesIO
import hashlib
import json
from pathlib import Path
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from docx import Document
from chatkeep import fetch_html, extract_stream, hydrate, find_conversation, parse_html
from word_export import build_docx

URL = 'https://chatgpt.com/share/6acdedca-cfa0-83ec-9622-dba167dfee84'


def run():
    html=fetch_html(URL)
    conversation=find_conversation(hydrate(extract_stream(html)))
    model=parse_html(html,URL)
    raw=conversation['linear_conversation']
    assert len(model['records'])==len(raw)
    compared=0
    for node,record in zip(raw,model['records']):
        assert record['node_id']==node.get('id')
        assert record['parent']==node.get('parent')
        source=node.get('message')
        if source is None:
            assert record['message'] is None
            continue
        normalized=record['message']
        assert normalized['id']==source['id']
        assert normalized['role']==source['author']['role']
        assert normalized['create_time']==source.get('create_time')
        assert normalized['content']==source.get('content')
        content=source['content']
        if isinstance(content.get('text'),str):expected=content['text']
        else:
            expected='\n'.join(p if isinstance(p,str) else p['text'] for p in content.get('parts',[]) if isinstance(p,str) or isinstance(p,dict) and isinstance(p.get('text'),str))
        assert normalized['text']==expected
        assert normalized['text_sha256']==hashlib.sha256(expected.encode()).hexdigest()
        compared+=1
    roles=Counter(n['message']['author']['role'] for n in raw if n.get('message'))
    assert dict(roles)==model['report']['roles']
    data=build_docx(model)
    doc=Document(BytesIO(data))
    bodies=[p.text for i,p in enumerate(doc.paragraphs) if i and doc.paragraphs[i-1].style.name=='Heading 2']
    assert bodies==[m['text'] for m in model['messages']]
    print(json.dumps({'http':'200 HTML', 'html_bytes':len(html.encode()),
                      'source_records_compared':compared,'report':model['report'],
                      'word_bytes':len(data),'word_messages_checked':len(bodies),
                      'source_and_word_checks':'passed', 'research_json_comparison':'尚未完成逐訊息比對'},ensure_ascii=False,indent=2))

if __name__=='__main__':run()
