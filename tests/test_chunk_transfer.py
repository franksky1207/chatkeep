"""Batch 5-1 Python protocol/API regressions; synthetic canonical packet only."""
from copy import deepcopy
import json
from io import BytesIO
import unittest
from unittest.mock import patch
from docx import Document
from chunk_transfer import sha, part_sha, validate_chunk, convert_chunk, CATEGORIES
from official_json import OfficialError
from test_official_json import call
import base64


def packet():
    messages=[{'id':'α😀','role':'user','text':'繁中\n"quote"\\slash😀'}, {'id':'β','role':'assistant','text':'English\t<&>'}]
    parts=[{'part_index':1,'first_message':1,'last_message':2,'message_count':2,'sequence_sha256':part_sha(messages)}]
    counts={'raw_node_count':2,'active_path_node_count':2,'active_path_message_count':2,'output_message_count':2,'off_branch_node_count':0,'skipped_categories':dict.fromkeys(CATEGORIES,0)}
    m={'schema_version':'chatkeep.transfer.v1','source':'official_json_derived','nonce':'0123456789abcdef0123456789abcdef','source_id':'source','title':'test','sequence_sha256':sha(parts),'parts':parts,'total_message_count':2,'total_text_chars':sum(len(x['text'])for x in messages),'source_counts':counts}
    m['conversation_key']=sha(m)
    return {'schema_version':'chatkeep.chunk.v1','manifest':m,'part_index':1,'messages':messages}


def encode(p):return json.dumps(p,ensure_ascii=False).encode()


class ChunkTests(unittest.TestCase):
    def test_real_api_and_explicit_source_scope(self):
        status,_,value=call(encode(packet()),PATH_INFO='/api/convert-chunk-v1')
        self.assertEqual(status,200);self.assertFalse(value['report']['complete']);self.assertTrue(value['report']['chunk_complete'])
        self.assertFalse(value['report']['source_verified_by_server'])
        self.assertEqual(value['transfer']['part_total'],1)
        doc=Document(BytesIO(base64.b64decode(value['document'])))
        self.assertEqual(sum(p.style.name=='Heading 2'for p in doc.paragraphs),2)

    def test_schema_order_counts_and_hashes_fail_closed(self):
        for mutate in (lambda p:p.update(complete=True),lambda p:p.update(part_index=True),
                       lambda p:p['messages'][0].update(text='CORRUPT'),lambda p:p['manifest']['parts'].reverse(),
                       lambda p:p['manifest']['source_counts'].update(output_message_count=False),
                       lambda p:p['manifest'].update(source='official_json'),lambda p:p['messages'][1].update(id='α😀')):
            p=packet();mutate(p)
            # reversing one part leaves the same manifest and is a valid no-op.
            if p==packet():continue
            with self.assertRaises(OfficialError):validate_chunk(encode(p))

    def test_too_many_parts_and_upload_bounds(self):
        p=packet();p['manifest']['parts']*=17
        with self.assertRaises(OfficialError) as e:validate_chunk(encode(p))
        self.assertEqual(e.exception.code,'PART_LIMIT')
        status,_,value=call(b'x',PATH_INFO='/api/convert-chunk-v1',CONTENT_LENGTH='2000001')
        self.assertEqual(status,413);self.assertFalse(value['ok'])

    def test_retry_after_and_shared_admission(self):
        from server import RateLimit
        with patch.object(RateLimit,'admit',return_value=False):
            status,headers,value=call(encode(packet()),PATH_INFO='/api/convert-chunk-v1')
        self.assertEqual(status,429);self.assertEqual(headers['Retry-After'],'60')
        self.assertEqual(value['error']['code'],'RATE_LIMITED')

    def test_timeout_worker_error_and_slot_recovery(self):
        with patch('json_api.CONVERSION_TIMEOUT',0):
            self.assertEqual(call(encode(packet()),PATH_INFO='/api/convert-chunk-v1')[2]['error']['code'],'CONVERSION_TIMEOUT')
        self.assertEqual(call(encode(packet()),PATH_INFO='/api/convert-chunk-v1')[0],200)

    def test_global_offset_and_source_binding(self):
        p=packet();parts=[dict(p['manifest']['parts'][0]),dict(p['manifest']['parts'][0],part_index=2,first_message=3,last_message=4)]
        p['part_index']=2;p['manifest'].update(parts=parts,total_message_count=4,total_text_chars=2*p['manifest']['total_text_chars'])
        p['manifest']['source_counts'].update(raw_node_count=4,active_path_node_count=4,active_path_message_count=4,output_message_count=4)
        p['manifest']['sequence_sha256']=sha(parts);del p['manifest']['conversation_key'];p['manifest']['conversation_key']=sha(p['manifest'])
        value=convert_chunk(encode(p));self.assertEqual(value['report']['volumes'][0]['first_message'],3)
        self.assertIn('_批002_',value['filename'])
        from docx.oxml.ns import qn
        doc=Document(BytesIO(base64.b64decode(value['document'])))
        names=[e.get(qn('w:name'))for e in doc.element.body.iter(qn('w:bookmarkStart'))]
        self.assertEqual(names,['ChatKeepMessage_3','ChatKeepMessage_4'])
