"""Stage 2-1: all fixtures synthetic, verify editable structures and text."""
import unittest
from io import BytesIO
from docx import Document
from semantic import parse, flatten
from word_export import build_docx

def model(text):
    return {"schema_version":"1.0","title":"合成元件驗收",
            "report":{"warnings":["測試樣本"]},
            "messages":[{"role":"assistant","text":text}]}

from docx_helpers import document_text as extract

class RichWordTests(unittest.TestCase):
    def test_box_row_nested_text_and_editability(self):
        source="<box>保留開頭 <text>正文甲</text> 末尾</box>\n<row><text>左側</text><text>右側</text></row>"
        doc=Document(BytesIO(build_docx(model(source),rich=True)))
        text=extract(doc)
        for fragment in ["保留開頭","正文甲","末尾","左側","右側"]:
            self.assertIn(fragment,text)
        self.assertGreaterEqual(len(doc.tables),2)
        self.assertNotIn("<box>",text)

    def test_markdown_table_and_heading(self):
        source="# 主題\n\n| 名稱 | 數值 |\n| --- | --- |\n| 甲 | 12 |\n\n- 第一點"
        doc=Document(BytesIO(build_docx(model(source),rich=True)))
        self.assertIn("主題",extract(doc))
        self.assertIn("甲",extract(doc))
        self.assertTrue(any(p.style.name=="List Bullet" for p in doc.paragraphs))
        self.assertGreaterEqual(len(doc.tables),1)

    def test_real_code_not_interpreted_as_ui(self):
        mark=chr(96)*3
        source=mark+"python\nprint('<box>不要刪掉</box>')\n"+mark
        doc=Document(BytesIO(build_docx(model(source),rich=True)))
        self.assertIn("<box>不要刪掉</box>",extract(doc))

    def test_chat_bubble(self):
        source="LINE｜阿哲：今晚有空嗎\nLINE｜我：有空"
        doc=Document(BytesIO(build_docx(model(source),rich=True)))
        self.assertIn("今晚有空嗎",extract(doc))
        self.assertIn("有空",extract(doc))
        self.assertGreaterEqual(len(doc.tables),2)

    def test_unknown_markup_not_dropped(self):
        source="<not-a-known-component>保留</not-a-known-component>"
        nodes,_=parse(source)
        self.assertIn("保留","".join(flatten(n) for n in nodes))
        doc=Document(BytesIO(build_docx(model(source),rich=True)))
        self.assertIn("保留",extract(doc))

if __name__=="__main__":
    unittest.main()
