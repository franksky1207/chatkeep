"""Batch 4 synthetic API, normalization, resource and editable DOCX tests."""
import base64
from contextlib import redirect_stdout, redirect_stderr
from io import BytesIO, StringIO
import json
import socket
import threading
import unittest
from unittest.mock import patch
from docx import Document
import json_api
from official_json import normalize_official, OfficialError, decode, MAX_INPUT
from official_worker import result
from server import Application, RateLimit
from docx_helpers import document_text


def msg(mid, role='user', text='中文 English 😀', **extra):
    return dict(id=mid, author={'role': role}, content={'content_type': 'text', 'parts': [text]}, **extra)


def conversation():
    # Deliberately shuffled mapping keys; only parent relations establish order.
    return {'title': '合成／測試😀', 'id': 'c1', 'current_node': 'a', 'mapping': {
        'other': {'parent': 'u', 'message': msg('other', 'assistant', 'OFF_BRANCH_SECRET')},
        'a': {'parent': 'u', 'message': msg('a', 'assistant', 'ANSWER')},
        'root': {'parent': None, 'message': None},
        'u': {'parent': 'root', 'message': msg('u', text='QUESTION')}}}


def body(c=None):
    return json.dumps(c if c is not None else conversation(), ensure_ascii=False).encode()


def call(data=None, **overrides):
    data = body() if data is None else data
    env = {'HTTP_HOST': 'localhost', 'PATH_INFO': '/api/convert-json', 'REQUEST_METHOD': 'POST',
           'CONTENT_TYPE': 'application/json', 'CONTENT_LENGTH': str(len(data)), 'wsgi.input': BytesIO(data),
           'REMOTE_ADDR': '127.0.0.1'}
    env.update(overrides)
    captured = []
    response = Application()(env, lambda status, headers: captured.append((status, dict(headers))))
    payload = json.loads(b''.join(response))
    return int(captured[0][0].split()[0]), captured[0][1], payload


class NormalizationTests(unittest.TestCase):
    def test_current_path_counts_and_order(self):
        model, report = normalize_official(body())
        self.assertEqual([m['id'] for m in model['messages']], ['u', 'a'])
        self.assertEqual([m['role'] for m in model['messages']], ['user', 'assistant'])
        self.assertEqual((report['raw_node_count'], report['active_path_node_count'],
                          report['active_path_message_count'], report['output_message_count']), (4, 3, 2, 2))
        self.assertEqual(report['off_branch_node_count'], 1)
        self.assertNotIn('SECRET', json.dumps(report))
        c = conversation(); c['current_node'] = 'other'
        self.assertEqual(normalize_official(body(c))[0]['messages'][-1]['id'], 'other')

    def test_hidden_tool_system_internal_empty(self):
        c = conversation(); parent = 'a'
        additions = [msg('hidden', metadata={'is_visually_hidden_from_conversation': True}),
                     msg('tool', 'tool'), msg('system', 'system'), msg('internal', 'assistant', channel='analysis'),
                     msg('call', 'assistant', recipient='python'), msg('empty', text='')]
        for m in additions:
            c['mapping'][m['id']] = {'parent': parent, 'message': m}; parent = m['id']
        c['current_node'] = parent
        model, report = normalize_official(body(c))
        self.assertEqual(len(model['messages']), 2)
        self.assertEqual(report['active_path_message_count'], 8)
        for cat in ('hidden', 'tool', 'system', 'internal', 'tool_call', 'empty'):
            self.assertEqual(report['skipped_categories'][cat], 1)
            self.assertIn('SKIPPED_' + cat.upper(), [w['code'] for w in report['warnings']])

    def test_invalid_branches_and_duplicate(self):
        for kind in ('missing_current', 'missing_parent', 'cycle', 'offcycle', 'id_mismatch', 'duplicate', 'children'):
            c = conversation()
            if kind == 'missing_current': del c['current_node']
            elif kind == 'missing_parent': c['mapping']['u']['parent'] = 'bad'
            elif kind == 'cycle': c['mapping']['root']['parent'] = 'a'
            elif kind == 'offcycle': c['mapping']['other']['parent'] = 'other'
            elif kind == 'id_mismatch': c['mapping']['a']['id'] = 'wrong'
            elif kind == 'duplicate': c['mapping']['a']['message']['id'] = 'u'
            elif kind == 'children': c['mapping']['root']['children'] = ['a']
            with self.subTest(kind=kind), self.assertRaises(OfficialError): normalize_official(body(c))

    def test_unsupported_and_unfinished_fail_closed(self):
        for content in ({'content_type': 'multimodal_text', 'parts': ['caption', {'image': 'x'}]},
                        {'content_type': 'future_type', 'parts': ['text']},
                        {'content_type': 'text', 'text': 'caption', 'parts': [{'image': 'secret'}]},
                        {'content_type': 'text', 'parts': [{'type': 'image', 'text': 'caption'}]}):
            c = conversation(); c['mapping']['a']['message']['content'] = content
            with self.assertRaises(OfficialError) as caught: normalize_official(body(c))
            self.assertEqual(caught.exception.code, 'UNSUPPORTED_CONTENT')
        c = conversation(); c['mapping']['a']['message']['status'] = 'in_progress'
        with self.assertRaises(OfficialError) as caught: normalize_official(body(c))
        self.assertEqual(caught.exception.code, 'UNFINISHED_MESSAGE')

    def test_schema_malformed_and_empty(self):
        for data in (b'{', b'\xff', b'{"a":1,"a":2}', b'{"x":NaN}', b'{"x":1e999}', b'[]', b'{}', body({'conversations': []})):
            with self.subTest(data=data), self.assertRaises(OfficialError): normalize_official(data)
        for field, value in [('title', {}), ('current_node', []), ('mapping', [] )]:
            c = conversation(); c[field] = value
            with self.assertRaises(OfficialError): normalize_official(body(c))
        c = conversation(); c['mapping']['u']['message'] = None; c['mapping']['a']['message'] = None
        with self.assertRaises(OfficialError): normalize_official(body(c))

    def test_limits_and_xml_characters(self):
        for data in (b' ' * (MAX_INPUT+1), b'['*33+b']'*33, b'{"x":"\\ud800"}'):
            with self.assertRaises(OfficialError): decode(data)
        with self.assertRaises(OfficialError): decode(json.dumps({'x':'\u0000'}).encode())
        with self.assertRaises(OfficialError): decode(b'[' + b'0,' * 100001 + b'0]')
        c = conversation(); c['mapping']['a']['message']['content']['parts'] = ['x'*1_000_001]
        with self.assertRaises(OfficialError) as caught: normalize_official(body(c))
        self.assertEqual(caught.exception.code, 'TEXT_LIMIT')
        with patch('official_json.MAX_NODES', 3), self.assertRaises(OfficialError): normalize_official(body())
        with patch('official_json.MAX_MESSAGES', 1), self.assertRaises(OfficialError): normalize_official(body())
        with patch('official_json.MAX_TEXT', 1), self.assertRaises(OfficialError): normalize_official(body())


class ExportTests(unittest.TestCase):
    def test_shared_renderer_and_docx_structure(self):
        text = '# 標題😀\n\n中文 **bold** *italic* `code`\n\n- 項目\n\n1. 編號\n\n```python\nprint("你好")\n```\n\n| A | B |\n|---|---|\n| 一 | 二 |\n\n[連結](https://example.com)'
        c = conversation(); c['mapping']['a']['message']['content']['parts'] = [text]
        packet = result(body(c)); data = base64.b64decode(packet['document'])
        doc = Document(BytesIO(data)); all_text = document_text(doc)
        for fragment in ('QUESTION', '標題😀', 'bold', 'italic', 'code', '你好', '一', '二'):
            self.assertIn(fragment, all_text)
        self.assertNotIn('OFF_BRANCH_SECRET', all_text)
        self.assertEqual(all_text.count('QUESTION'), 1)
        self.assertGreaterEqual(len(doc.tables), 1)
        self.assertTrue(any(r.reltype.endswith('/hyperlink') for r in doc.part.rels.values()))
        xml = doc.part.blob.decode()
        self.assertIn('<w:b', xml); self.assertIn('<w:i', xml)
        self.assertNotIn('/', packet['filename'])
        from word_export import WordVolume
        with patch('conversation_export.WordVolume', wraps=WordVolume) as shared:
            result(body()); self.assertTrue(shared.called)

    def test_malicious_html_no_fetch_or_execution(self):
        c = conversation(); c['title'] = '../CON/\\?'
        c['mapping']['a']['message']['content']['parts'] = ['<script>PRIVATE_SENTINEL</script>\n<img src="http://127.0.0.1/private">\n[bad](javascript:alert(1))']
        with patch('urllib.request.urlopen', side_effect=AssertionError('network forbidden')):
            packet = result(body(c))
        doc = Document(BytesIO(base64.b64decode(packet['document'])))
        self.assertIn('PRIVATE_SENTINEL', document_text(doc))
        self.assertFalse(any(r.target_ref.startswith('javascript:') for r in doc.part.rels.values()))
        self.assertNotIn('PRIVATE_SENTINEL', json.dumps(packet['report']))


class APITests(unittest.TestCase):
    def test_success_real_subprocess_and_privacy(self):
        out, err = StringIO(), StringIO()
        with redirect_stdout(out), redirect_stderr(err): status, headers, packet = call()
        self.assertEqual(status, 200)
        self.assertTrue(packet['ok']); self.assertEqual(headers['Cache-Control'], 'no-store')
        self.assertNotIn('Access-Control-Allow-Origin', headers)
        self.assertEqual(out.getvalue()+err.getvalue(), '')
        doc = Document(BytesIO(base64.b64decode(packet['document'])))
        self.assertIn('ANSWER', document_text(doc))
        self.assertNotIn('OFF_BRANCH_SECRET', document_text(doc))

    def test_error_contract(self):
        cases = [(b'{', {}, 400, 'INVALID_JSON'), (b'[]', {}, 422, 'SINGLE_CONVERSATION_REQUIRED'),
                 (body(), {'CONTENT_TYPE':'application/zip'}, 415, 'UNSUPPORTED_MEDIA_TYPE'),
                 (body(), {'CONTENT_LENGTH':'2000001'}, 413, 'INPUT_TOO_LARGE'),
                 (body(), {'HTTP_ORIGIN':'https://evil.test'}, 403, 'ORIGIN_FORBIDDEN'),
                 (body(), {'REQUEST_METHOD':'OPTIONS'}, 405, 'METHOD_NOT_ALLOWED'),
                 (body(), {'CONTENT_LENGTH':'bad'}, 400, 'INVALID_LENGTH'),
                 (body(), {'HTTP_TRANSFER_ENCODING':'chunked'}, 415, 'UNSUPPORTED_ENCODING'),
                 (body(), {'HTTP_CONTENT_ENCODING':'gzip'}, 415, 'UNSUPPORTED_ENCODING'),
                 (body(), {'CONTENT_LENGTH':str(len(body())+1)}, 400, 'INCOMPLETE_BODY')]
        for data, env, expected, code in cases:
            with self.subTest(code=code):
                status, _, packet = call(data, **env)
                self.assertEqual(status, expected); self.assertEqual(packet['error']['code'], code)
                self.assertNotIn('document', packet)

    def test_admission_and_slot_release(self):
        with patch('server.SLOTS', threading.BoundedSemaphore(0)):
            self.assertEqual(call()[2]['error']['code'], 'BUSY')
        with patch.object(RateLimit, 'admit', return_value=False):
            self.assertEqual(call()[2]['error']['code'], 'RATE_LIMITED')
        slot = threading.BoundedSemaphore(1)
        with patch('server.SLOTS', slot):
            self.assertEqual(call(b'{')[0], 400)
            self.assertTrue(slot.acquire(False)); slot.release()

    def test_timeout_kills_and_reaps(self):
        slot = threading.BoundedSemaphore(1)
        with patch('server.SLOTS', slot), patch('json_api.CONVERSION_TIMEOUT', 0):
            self.assertEqual(call()[2]['error']['code'], 'CONVERSION_TIMEOUT')
        self.assertTrue(slot.acquire(False)); slot.release()

    def test_slow_upload_is_actually_aborted(self):
        reader, writer = socket.socketpair()
        stream = reader.makefile('rb')
        try:
            with patch('json_api.UPLOAD_TIMEOUT', .02):
                status, _, packet = call(**{'wsgi.input': stream, 'gunicorn.socket': reader})
            self.assertEqual(status, 408)
            self.assertEqual(packet['error']['code'], 'UPLOAD_TIMEOUT')
        finally:
            stream.close(); reader.close(); writer.close()

    def test_worker_crash_safe_error(self):
        with patch('json_api.subprocess.Popen') as p:
            p.return_value.communicate.return_value = (b'', None)
            p.return_value.returncode = -9
            p.return_value.poll.return_value = -9
            self.assertEqual(call()[2]['error']['code'], 'RESOURCE_LIMIT')

    def test_upload_timeout(self):
        with patch('json_api.subprocess.Popen'):
            class Slow:
                def read(self, size): raise TimeoutError('PRIVATE')
            status, _, packet = call(**{'wsgi.input': Slow()})
            self.assertEqual(status, 408); self.assertNotIn('PRIVATE', json.dumps(packet))


if __name__ == '__main__': unittest.main()
