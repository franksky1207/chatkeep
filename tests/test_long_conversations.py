"""Batch 5: measured structural capacity, transactional volumes, both API sources."""
import base64
from io import BytesIO
import json
from pathlib import Path
import subprocess
import sys
import threading
import unittest
from unittest.mock import patch
from zipfile import ZipFile
from docx import Document
from docx.oxml.ns import qn
from conversation_export import export_conversation, envelope, verify_docx, digest
from export_limits import Capacity, CAPACITY, MAX_ARTIFACT
from official_worker import result
from word_export import ExportError, LimitedBuffer, build_docx
from test_official_json import conversation, body, call
from test_integration import html_for
from chatkeep import parse_html
from docx_helpers import document_text
import server


def model(texts):
    return {'schema_version':'1.0','title':'../文明戰線101/😀',
            'report':{'warnings':[]}, 'messages':[{'id':'m%d'%i, 'role':'user' if i%2==0 else 'assistant',
                'text':t} for i,t in enumerate(texts)]}


def payloads(artifact):
    if artifact['artifact_type']=='docx': return [artifact['data']]
    with ZipFile(BytesIO(artifact['data'])) as archive:
        assert archive.testzip() is None
        return [archive.read(n) for n in archive.namelist()]


def positions(data):
    doc=Document(BytesIO(data))
    return [n.get(qn('w:name')) for n in doc.element.body.iter(qn('w:bookmarkStart'))]


class CapacityTests(unittest.TestCase):
    def test_short_and_medium_prefer_single(self):
        for texts in (['one','two'], [('中文abcd'*5000) for _ in range(20)], [('line\n'*1000) for _ in range(20)]):
            artifact=export_conversation(model(texts),'share_link')
            self.assertEqual(artifact['artifact_type'],'docx')
            self.assertEqual(artifact['report']['docx_volume_count'],1)
            self.assertTrue(artifact['report']['complete'])

    def test_old_100000_counter_safe_literal(self):
        texts=['```text\n'+('<'*10000)+'\n```' for _ in range(20)]
        self.assertGreater(sum(t.count('\n')+t.count('<') for t in texts),100000)
        m=model(texts)
        data=build_docx(m)
        self.assertEqual(len(positions(data)),20)
        self.assertEqual(document_text(Document(BytesIO(data))).count('<'),200000)
        artifact=export_conversation(m,'share_link')
        self.assertEqual(artifact['artifact_type'],'docx')

    def test_real_structure_rollover_is_contiguous(self):
        texts=['TOKEN_%02d\n'%i + '正文\n'*30 for i in range(35)]
        m=model(texts)
        # Small measured budgets exercise every boundary cheaply; no fixed message split.
        artifact=export_conversation(m,'share_link',capacity=Capacity(800,8_000_000,5000))
        self.assertEqual(artifact['artifact_type'],'zip')
        parts=artifact['report']['volumes']; docs=payloads(artifact)
        self.assertEqual(len(parts),len(docs))
        seen=[]; output=''
        for part,data in zip(parts,docs):
            expected=['ChatKeepMessage_%d'%n for n in range(part['first_message'],part['last_message']+1)]
            self.assertEqual(positions(data),expected); seen.extend(expected)
            self.assertEqual(part['message_count'],len(expected))
            self.assertEqual(part['sequence_sha256'],digest(m['messages'][part['first_message']-1:part['last_message']]))
            output+=document_text(Document(BytesIO(data)))
            self.assertIn('_第%d卷.docx'%part['number'],part['filename'])
            self.assertNotIn('/',part['filename'])
        self.assertEqual(seen,['ChatKeepMessage_%d'%n for n in range(1,36)])
        for i in range(35): self.assertEqual(output.count('TOKEN_%02d'%i),1)
        self.assertEqual(sorted(range(35),key=lambda i:output.index('TOKEN_%02d'%i)),list(range(35)))
        report=artifact['report']
        self.assertEqual(report['output_message_count'],report['expected_output_message_count'])
        for key in ('truncated','duplicated','missing'): self.assertFalse(report[key])
        self.assertIsNone(report['raw_node_count'])
        packet=envelope(artifact)
        self.assertEqual(packet['artifact_type'],'zip'); self.assertEqual(packet['mime_type'],'application/zip')
        self.assertEqual(base64.b64decode(packet['document']),artifact['data'])

    def test_large_real_tables_and_code(self):
        table='|A|B|C|\n|---|---|---|\n'+'|甲|乙|丙|\n'*200
        text='```python\nprint("你好")\n```\n'*100
        artifact=export_conversation(model([table]*3+[text]*4),'share_link')
        self.assertEqual(artifact['artifact_type'],'docx')
        doc=Document(BytesIO(artifact['data']))
        self.assertEqual(len(doc.tables),3)
        self.assertEqual(len(doc.tables[0].rows),201)
        self.assertEqual(document_text(doc).count('甲'),600)
        self.assertEqual(document_text(doc).count('print("你好")'),400)

    def test_table_size_and_single_message_fail_without_cutting(self):
        for text in ('```text\n'+'line\n'*80000+'```', 'x'*1_000_001,
                     '|A|B|C|\n|---|---|---|\n'+'|a|b|c|\n'*2000):
            with self.assertRaises(ExportError) as caught: export_conversation(model([text]),'share_link')
            self.assertEqual(caught.exception.code,'MESSAGE_TOO_LARGE')

    def test_same_single_message_can_use_larger_safe_prose(self):
        artifact=export_conversation(model(['中文'*200000]),'share_link')
        self.assertEqual(artifact['artifact_type'],'docx')
        self.assertIn('中文'*100,document_text(Document(BytesIO(artifact['data']))))

    def test_volume_limit_and_response_budget(self):
        with patch('conversation_export.MAX_VOLUMES',1), self.assertRaises(ExportError) as caught:
            export_conversation(model(['x\n'*100]*5),'share_link',capacity=Capacity(500,8_000_000,5000))
        self.assertEqual(caught.exception.code,'VOLUME_LIMIT')
        with patch('word_export.MAX_ARTIFACT',100), self.assertRaises(ExportError) as caught:
            export_conversation(model(['short']),'share_link')
        self.assertEqual(caught.exception.code,'OUTPUT_LIMIT')
        buffer=LimitedBuffer()
        with self.assertRaises(ExportError): buffer.write(b'x'*(MAX_ARTIFACT+1))

    def test_pack_failure_returns_no_partial_zip_and_no_tempfile(self):
        from word_export import WordVolume
        original=WordVolume.pack
        def pack(volume,number,*args):
            if number==2: raise ExportError('synthetic failure','OUTPUT_LIMIT')
            return original(volume,number,*args)
        with (patch('tempfile.NamedTemporaryFile',side_effect=AssertionError('private disk use')),
              patch.object(WordVolume,'pack',pack)):
            with self.assertRaises(ExportError):
                export_conversation(model(['x\n'*100]*8),'share_link',capacity=Capacity(500,8_000_000,5000))

    def test_cancel_during_later_volume_resets_renderer_context(self):
        from export_limits import OBSERVER
        from word_export import WordVolume
        original=WordVolume.pack
        cancel=threading.Event()
        def pack(*args,**kwargs):
            data=original(*args,**kwargs); cancel.set(); return data
        def check():
            if cancel.is_set(): raise RuntimeError('cancelled')
        with patch.object(WordVolume,'pack',pack),self.assertRaises(RuntimeError):
            export_conversation(model(['x\n'*100]*8),'share_link',capacity=Capacity(500,8_000_000,5000),check=check)
        self.assertIsNone(OBSERVER.get())
        self.assertEqual(export_conversation(model(['next']),'share_link')['artifact_type'],'docx')

    def test_serialized_ledger_and_table_compatibility(self):
        artifact=export_conversation(model(['<box><row><col><box>A</box></col><col><box>B</box></col></row></box>']*3),'share_link')
        data=artifact['data']; verify_docx(data,1,3)
        with self.assertRaises(ExportError): verify_docx(data,1,2)
        doc=Document(BytesIO(data))
        self.assertTrue(list(doc.element.body.iter(qn('w:tc'))))
        for cell in doc.element.body.iter(qn('w:tc')): self.assertEqual(cell[-1].tag,qn('w:p'))


class SourceTests(unittest.TestCase):
    def test_official_active_branch_and_multipart(self):
        c=conversation(); parent='a'
        for i in range(7):
            key='tail%d'%i
            c['mapping'][key]={'parent':parent,'message':{'id':key,'author':{'role':'assistant'},
                'content':{'content_type':'text','parts':['tail%d\n'%i+'x\n'*100]}}}
            parent=key
        c['current_node']=parent
        from official_json import normalize_official
        m,report=normalize_official(body(c))
        artifact=export_conversation(m,'official_json',report,capacity=Capacity(500,8_000_000,5000))
        self.assertEqual(artifact['artifact_type'],'zip')
        text=''.join(document_text(Document(BytesIO(d))) for d in payloads(artifact))
        self.assertNotIn('OFF_BRANCH_SECRET',text)
        self.assertEqual(artifact['report']['raw_node_count'],11)
        self.assertEqual(artifact['report']['active_path_message_count'],9)
        self.assertEqual(artifact['report']['output_message_count'],9)

    def test_both_worker_sources_use_same_core(self):
        packet=result(body())
        self.assertEqual(packet['api_version'],2)
        m=parse_html(html_for())
        with patch('chatkeep.extract_url',return_value=m):
            share=result(json.dumps({'url':'https://chatgpt.com/share/11111111-1111-1111-1111-111111111111'}).encode(),'share')
        self.assertEqual(packet['artifact_type'],share['artifact_type'])
        self.assertEqual(share['report']['source'],'share_link')
        self.assertIsNone(share['report']['raw_node_count'])
        self.assertEqual(share['report']['expected_output_message_count'],2)
        for packet in (packet,share):
            self.assertTrue(packet['report']['complete']);self.assertEqual(packet['report']['docx_volume_count'],1)

    def test_shared_legacy_report_and_multipart_error(self):
        events=[]; m=parse_html(html_for())
        with patch('server.extract_url',return_value=m): server.convert('https://chatgpt.com/share/11111111-1111-1111-1111-111111111111',events.append)
        self.assertEqual(events[-1]['stage'],'done')
        self.assertTrue(events[-1]['report']['export']['complete'])
        with self.assertRaises(ExportError) as caught:
            export_conversation(model(['x\n'*100]*4),'share_link',capacity=Capacity(500,8_000_000,5000),allow_zip=False)
        self.assertEqual(caught.exception.code,'MULTIPART_REQUIRED')

    def test_v2_does_not_silently_drop_ambiguous_content(self):
        m=model(['caption'])
        m['messages'][0]['content']={'content_type':'text','text':'caption','parts':[{'image':'not-exportable'}]}
        with self.assertRaises(ExportError) as caught: export_conversation(m,'share_link')
        self.assertEqual(caught.exception.code,'UNSUPPORTED_CONTENT')

    def test_v2_transport_validation_and_worker_mode(self):
        status,_,packet=call(b'{"url":"https://evil.test"}',PATH_INFO='/api/convert-v2')
        self.assertEqual(status,400); self.assertEqual(packet['error']['code'],'INVALID_URL')
        self.assertEqual(call(b'{', PATH_INFO='/api/convert-v2')[2]['error']['code'], 'INVALID_JSON')
        packet=result(body())
        with patch('json_api.subprocess.Popen') as p:
            p.return_value.communicate.return_value=(json.dumps({'status':200,'value':packet}).encode(),None)
            p.return_value.returncode=0; p.return_value.poll.return_value=0
            request=b'{"url":"https://chatgpt.com/share/11111111-1111-1111-1111-111111111111"}'
            self.assertEqual(call(request,PATH_INFO='/api/convert-v2')[0],200)
            self.assertEqual(p.call_args.args[0][-1],'share')

    def test_timeout_reaps_actual_process_and_closes_pipes(self):
        real_popen = subprocess.Popen
        children = []
        def launch(*args, **kwargs):
            process = real_popen(*args, **kwargs); children.append(process); return process
        with patch('json_api.subprocess.Popen', side_effect=launch), patch('json_api.CONVERSION_TIMEOUT', 0):
            self.assertEqual(call()[2]['error']['code'], 'CONVERSION_TIMEOUT')
        self.assertTrue(children)
        for process in children:
            self.assertIsNotNone(process.poll())
            self.assertTrue(process.stdin.closed and process.stdout.closed)

    def test_legacy_disconnect_checks_actual_exporter(self):
        slots = threading.BoundedSemaphore(1)
        slots.acquire()
        m = model(['line\n'*100]*100)
        with patch('server.SLOTS', slots), patch('server.extract_url', return_value=m):
            stream = server.Application().stream('https://chatgpt.com/share/11111111-1111-1111-1111-111111111111')
            next(stream); stream.close()
            self.assertTrue(slots.acquire(timeout=2))
            slots.release()

    def test_real_worker_memory_failure_is_safe(self):
        command='import official_worker as w; w.result=lambda *args:bytearray(512*1024**2); w.main()'
        p=subprocess.run([sys.executable,'-c',command],input=b'{}',stdout=subprocess.PIPE,stderr=subprocess.PIPE,timeout=10)
        value=json.loads(p.stdout)
        self.assertEqual(value['status'],503)
        self.assertEqual(value['value']['error']['code'],'RESOURCE_LIMIT')
        self.assertNotIn('document',value['value']); self.assertEqual(p.stderr,b'')

    def test_real_cpu_limit_terminates_worker(self):
        command = """import official_worker as w, resource
original = resource.setrlimit
def limits(kind, value): original(kind, (1, 1) if kind == resource.RLIMIT_CPU else value)
w.resource.setrlimit = limits
def spin(*args):
    while True: pass
w.result = spin
w.main()
"""
        p=subprocess.run([sys.executable,'-c',command],input=b'{}',stdout=subprocess.PIPE,stderr=subprocess.PIPE,timeout=5)
        self.assertIn(p.returncode,(-9,-24))
        self.assertEqual(p.stdout+p.stderr,b'')


if __name__=='__main__': unittest.main()
