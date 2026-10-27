"""Stage 5-1 regression contracts. Synthetic content only."""
from io import BytesIO
import json
import socket
import threading
import time
import unittest
from unittest.mock import patch, Mock
from docx import Document
from docx.oxml.ns import qn
import chatkeep
import server
from semantic import parse, flatten
from word_export import build_docx, ExportError, filename_for
from test_integration import fixture, message, html_for, URL, FakeResponse
from docx_helpers import document_text


def model(text):
    return chatkeep.normalize(fixture([message('synthetic', 'assistant', text)]))


class ContentAcceptance(unittest.TestCase):
    def test_nested_native_structures_no_loss_or_duplicate(self):
        text='<box>開頭唯一<row><col><text>左欄唯一</text><card><caption>巢狀唯一</caption></card></col><col><text>右欄唯一</text></col></row>結尾唯一</box>'
        doc=Document(BytesIO(build_docx(model(text))))
        output=document_text(doc)
        fragments=['開頭唯一','左欄唯一','巢狀唯一','右欄唯一','結尾唯一']
        for token in fragments:self.assertEqual(output.count(token),1)
        self.assertEqual(sorted(fragments,key=output.index),fragments)
        self.assertGreaterEqual(len(list(doc.element.body.iter(qn('w:tbl')))),3)
        self.assertNotIn('<box>',output)

    def test_literal_code_fence_length_and_inline_markup(self):
        source="````python\n<box>程式甲</box>\n```\n程式乙\n````\n行內 `<text>程式丙</text>` 保留"
        doc=Document(BytesIO(build_docx(model(source))))
        output=document_text(doc)
        self.assertIn("<box>程式甲</box>\n```\n程式乙",'\n'.join(p.text for p in doc.paragraphs))
        self.assertIn('`<text>程式丙</text>`',output)

    def test_malformed_dynamic_unknown_and_decorative_labels(self):
        for source, tokens in [('<box>未閉合正文',['未閉合正文']),
                               ('<box label={dynamic}>動態正文</box>',['動態正文','{dynamic}']),
                               ('<?instruction?><unknown>未知正文</unknown>',['未知正文','<?instruction?>']),
                               ('<button label="選項甲"/><icon alt="重要標籤"/><text>正文乙</text>', ['選項甲','重要標籤','正文乙'])]:
            with self.subTest(source=source):
                doc=Document(BytesIO(build_docx(model(source))))
                for token in tokens:self.assertIn(token,document_text(doc))

    def test_one_column_row_keeps_children_in_flow_not_one_indented_cell(self):
        source='<box><row>\n  <box>\n    <text>正文甲</text>\n    <caption>正文乙</caption>\n  </box>\n</row></box>'
        doc=Document(BytesIO(build_docx(model(source))))
        self.assertEqual(document_text(doc).count('正文甲'),1)
        self.assertEqual(document_text(doc).count('正文乙'),1)
        self.assertEqual(len(list(doc.element.body.iter(qn('w:br')))),0)

    def test_self_closing_text_and_box_values_are_not_empty(self):
        doc=Document(BytesIO(build_docx(model('<text value="文字甲"/><box value="文字乙"/>'))))
        self.assertIn('文字甲',document_text(doc))
        self.assertIn('文字乙',document_text(doc))

    def test_jsx_layout_attributes_do_not_swallow_body_or_execute(self):
        source='<box pad={12} onClick={()=>alert("<ignored>")}><row><text>正文甲</text><text>正文乙</text></row></box>'
        doc=Document(BytesIO(build_docx(model(source))))
        output=document_text(doc)
        self.assertEqual(output.count('正文甲'),1)
        self.assertEqual(output.count('正文乙'),1)
        self.assertNotIn('onClick',output)
        self.assertGreaterEqual(len(list(doc.element.body.iter(qn('w:tbl')))),2)
        literal='行內 `<box pad={12}>真正程式</box>`'
        self.assertIn('`<box pad={12}>真正程式</box>`',document_text(Document(BytesIO(build_docx(model(literal))))))

    def test_markdown_table_after_prose_and_escaped_pipe(self):
        source='前言\n| 名稱 | 值 |\n| --- | --- |\n| A\\|B | `x|y` |\n> 引用正文\n- 項目正文\n\\(a+b\\)'
        doc=Document(BytesIO(build_docx(model(source))))
        self.assertEqual(len(doc.tables[0].columns),2)
        self.assertEqual(doc.tables[0].cell(1,0).text,'A|B')
        self.assertEqual(doc.tables[0].cell(1,1).text,'`x|y`')
        for text in ['前言','引用正文','項目正文','\\(a+b\\)']:
            self.assertIn(text,document_text(doc))

    def test_line_bubbles_names_time_side_without_prose_misclassification(self):
        doc=Document(BytesIO(build_docx(model('說明：不能變成泡泡\nLINE｜阿明：19:30 訊息甲\nLINE｜我：19:31 訊息乙'))))
        self.assertEqual(len(doc.tables),2)
        self.assertIn('阿明 · 19:30',doc.tables[0].cell(0,0).text)
        self.assertIn('我 · 19:31',doc.tables[1].cell(0,1).text)
        self.assertIn('說明：不能變成泡泡',[p.text for p in doc.paragraphs])

    def test_explicit_list_numbers_and_table_repeating_header(self):
        doc=Document(BytesIO(build_docx(model('5. 項目甲\n9. 項目乙\n| 欄一 | 欄二 |\n| --- | --- |\n| 值甲 | 值乙 |'))))
        self.assertIn('5. 項目甲',document_text(doc))
        self.assertIn('9. 項目乙',document_text(doc))
        self.assertIn('tblHeader',doc.tables[0].rows[0]._tr.xml)

    def test_image_and_parse_warnings_reach_document_and_api(self):
        source='<box><AsyncImage query="合成圖片"/></box>\n<box>未閉合'
        warnings=[]
        doc=Document(BytesIO(build_docx(model(source),warnings_out=warnings)))
        self.assertTrue(warnings)
        self.assertIn('[圖片未嵌入]',document_text(doc))
        self.assertIn('轉換警告',document_text(doc))

    def test_long_card_can_split_and_nested_layout_polished(self):
        doc=Document(BytesIO(build_docx(model('<box><text>'+('長篇正文\n'*1000)+'</text></box>'))))
        self.assertNotIn('cantSplit',doc.tables[0].rows[0]._tr.xml)
        self.assertIn('w:val="top"',doc.tables[0].cell(0,0)._tc.xml)
        self.assertIn('長篇正文',document_text(doc))
        self.assertAlmostEqual(doc.sections[0].page_width.cm,21,places=1)
        self.assertIn('PAGE',doc.sections[0].footer._element.xml)

    def test_limits_fail_explicitly_instead_of_truncation(self):
        data=model('超長')
        data['messages']*=5001
        with self.assertRaises(ExportError):build_docx(data)
        data=model('正文\n'*100001)
        with self.assertRaises(ExportError):build_docx(data)

    def test_reserved_names_with_extensions_and_unicode(self):
        for value in ['CON','con.txt','LPT1.any','COM¹']:
            self.assertTrue(filename_for(value).startswith('_'))
        self.assertEqual(filename_for('中文😀'),'中文😀.docx')


class ProgressAcceptance(unittest.TestCase):
    def test_processed_nodes_not_started_nodes(self):
        events=[]
        with self.assertRaises(chatkeep.ExtractionError):
            chatkeep.normalize({'linear_conversation':['invalid']},on_progress=lambda c,t:events.append((c,t)))
        self.assertEqual(events,[(0,1)])
        events=[]
        chatkeep.normalize(fixture(),on_progress=lambda c,t:events.append((c,t)))
        self.assertEqual(events,[(0,2),(2,2)])

    def test_word_counts_and_pack_before_actual_save(self):
        events=[]
        from docx.document import Document as DocumentClass
        original=DocumentClass.save
        def save(doc,target):events.append('save');return original(doc,target)
        with patch.object(DocumentClass,'save',save):
            build_docx(model('正文'),on_progress=lambda c,t:events.append((c,t)),on_pack=lambda:events.append('pack'))
        self.assertEqual(events,[(0,1),(1,1),'pack','save'])

    def test_unknown_download_total_still_reports_real_bytes(self):
        events=[]
        response=FakeResponse(b'x'*150000)
        with patch('chatkeep.urllib.request.build_opener') as opener:
            opener.return_value.open.return_value=response
            chatkeep.fetch_html(URL,on_progress=lambda c,t:events.append((c,t)))
        self.assertEqual(events,[(65536,None),(131072,None),(150000,None)])

    def test_http_stage_counters_and_rendering_warnings(self):
        events=[]
        with patch('chatkeep.fetch_html',return_value=html_for(fixture([message('a',text='<AsyncImage query="test"/>')]))):
            server.convert(URL,events.append)
        stages=[e['stage'] for e in events]
        self.assertLess(stages.index('word'),stages.index('pack'))
        self.assertEqual(stages[-1],'done')
        self.assertIsNone(next(e for e in events if e['stage']=='pack')['percent'])
        for event in events:
            if event.get('total'):
                self.assertEqual(event['percent'],int(event['current']*100/event['total']))
        self.assertTrue(events[-1]['report']['rendering_warnings'])


class SafetyAcceptance(unittest.TestCase):
    def test_direct_dns_private_or_mixed_never_connects(self):
        for values in [['127.0.0.1'],['::1'],['169.254.169.254'],['10.0.0.1'],['104.18.32.7','192.168.1.1']]:
            results=[(socket.AF_INET6 if ':' in ip else socket.AF_INET,socket.SOCK_STREAM,6,'',(ip,443)) for ip in values]
            with patch('chatkeep.socket.getaddrinfo',return_value=results),patch('chatkeep.socket.socket') as connect:
                with self.assertRaises(chatkeep.ExtractionError):chatkeep.public_connection(('chatgpt.com',443))
                connect.assert_not_called()

    def test_direct_dns_pins_numeric_address_once(self):
        with patch('chatkeep.socket.getaddrinfo',return_value=[(socket.AF_INET,socket.SOCK_STREAM,6,'',('104.18.32.7',443))]) as dns,patch('chatkeep.socket.socket') as factory:
            chatkeep.public_connection(('chatgpt.com',443))
            factory.return_value.connect.assert_called_once_with(('104.18.32.7',443))
            dns.assert_called_once()

    def request(self,payload,**extra):
        body=json.dumps(payload).encode();status=[]
        env={'REQUEST_METHOD':'POST','PATH_INFO':'/api/convert','HTTP_HOST':'localhost','CONTENT_TYPE':'application/json','CONTENT_LENGTH':str(len(body)),'wsgi.input':BytesIO(body),'REMOTE_ADDR':'test'}
        env.update(extra)
        result=server.Application()(env,lambda code,headers:status.append(code))
        return status[0],result

    def test_bad_host_origin_body_and_cross_site(self):
        for extra in [{'HTTP_HOST':'evil.test'},{'HTTP_ORIGIN':'http://evil.test'},{'HTTP_ORIGIN':'http://['},{'HTTP_SEC_FETCH_SITE':'cross-site'},{'HTTP_TRANSFER_ENCODING':'chunked'}]:
            code,_=self.request({'url':URL},**extra)
            self.assertIn(code[:3],['400','403'])
        code,_=self.request({'url':URL},CONTENT_LENGTH='999999')
        self.assertTrue(code.startswith('413'))
        for payload in [None,[],{'url':{}},{'url':'https://127.0.0.1'}]:
            code,_=self.request(payload);self.assertTrue(code.startswith('400'))

    def test_rate_storage_and_admission_bounded(self):
        rates=server.RateLimit()
        self.assertTrue(all(rates.admit('one') for _ in range(10)))
        self.assertFalse(rates.admit('one'))
        with patch('server.time.monotonic',return_value=time.monotonic()+61):
            self.assertTrue(rates.admit('one'))

    def test_disconnect_stops_worker_and_releases_slot(self):
        ended=threading.Event()
        def work(url,emit):
            try:
                for i in range(10000):emit({'stage':'parse','message':'synthetic','percent':None})
            finally:ended.set()
        app=server.Application()
        with patch('server.convert',side_effect=work):
            code,body=self.request({'url':URL})
            self.assertTrue(code.startswith('200'))
            next(body);body.close()
            self.assertTrue(ended.wait(2))
        acquired=[server.SLOTS.acquire(blocking=False) for _ in range(2)]
        self.assertEqual(acquired,[True,True])
        for _ in acquired:server.SLOTS.release()

    def test_busy_conversion_has_no_success_body(self):
        self.assertTrue(server.SLOTS.acquire());self.assertTrue(server.SLOTS.acquire())
        try:
            code,body=self.request({'url':URL})
            self.assertTrue(code.startswith('429'))
            self.assertNotIn(b'document',b''.join(body))
        finally:server.SLOTS.release();server.SLOTS.release()
