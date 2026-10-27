"""Batch 3-1 layout contract tests; run together with batch 5-1 CI."""
import unittest
from io import BytesIO
from docx import Document
from docx.oxml.ns import qn
from word_export import build_docx

def fixture():
    return {"schema_version": "1.0", "title": "合成長篇文件",
            "report": {"warnings": []},
            "messages": [
                {"role": "user", "text": "# 章節\n\n第一段\n第二行"},
                {"role": "assistant", "text": "<box><text>完整方框</text></box>"},
                {"role": "assistant", "text": "LINE｜我：測試訊息"},
            ]}

class WordLayoutTests(unittest.TestCase):
    def test_reopen_page_and_heading_styles(self):
        doc = Document(BytesIO(build_docx(fixture())))
        self.assertEqual(doc.paragraphs[0].text, "合成長篇文件")
        self.assertEqual(doc.paragraphs[0].style.name, "Title")
        self.assertTrue(any(p.style.name == "Heading 2" for p in doc.paragraphs))
        self.assertIn("完整方框", "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells))
        normal = doc.styles["Normal"]
        self.assertIsNotNone(normal.font.size)
        self.assertTrue(doc.sections[0].footer.paragraphs)
        footer_xml = doc.sections[0].footer._element.xml
        self.assertIn("PAGE", footer_xml)
        self.assertIn("cantSplit", doc.tables[0]._tbl.xml)

    def test_large_conversation_no_page_per_message(self):
        source = fixture()
        source["messages"] = [{"role": "user" if i % 2 else "assistant", "text": "內容 " + str(i)} for i in range(600)]
        doc = Document(BytesIO(build_docx(source)))
        self.assertEqual(sum(p.style.name == "Heading 2" for p in doc.paragraphs), 600)
        self.assertNotIn('w:type="page"', doc._body._body.xml)

if __name__ == "__main__":
    unittest.main()
