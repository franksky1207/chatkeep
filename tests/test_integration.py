"""All fixtures are synthetic. No live conversation files in the repository."""
from collections import Counter
from io import BytesIO
import hashlib
import json
import socket
import threading
import unittest
import urllib.error
import urllib.request
from unittest.mock import patch
from docx import Document
import chatkeep
from chatkeep import ExtractionError, hydrate, normalize, parse_html, fetch_html
from server import Server, Handler
from word_export import ExportError, build_docx, filename_for
from test_chatkeep import Encoder
from docx_helpers import document_text

URL = 'https://chatgpt.com/share/11111111-1111-1111-1111-111111111111'
TEXT = '# 標題\n\n繁體中文，標點！😀\t<&>\\跳脫\n```python\nprint("你好")\n```\n<box><row><text>正文不能刪</text></row></box>\n'


def message(mid, role='user', text=TEXT, **extra):
    return {'id': mid, 'author': {'role': role}, 'content': {'content_type': 'text', 'parts': [text]},
            'create_time': 1234567890.5, 'status': 'finished_successfully', **extra}


def fixture(messages=None):
    messages = messages or [message('a'), message('b', 'assistant')]
    return {'title': '合成測試／Word', 'linear_conversation': [{'message': m} for m in messages]}


def html_for(conv=None):
    encoder = Encoder(); encoder.add({'loaderData': {'route': conv or fixture()}})
    raw = json.dumps(encoder.flat, ensure_ascii=False)
    # Chunk inside an escape or Unicode sequence must round-trip without eval.
    parts = [raw[:17], raw[17:105], raw[105:]]
    return '<script>' + ''.join('window.__reactRouterContext.streamController.enqueue(' + json.dumps(p) + ');' for p in parts) + '</script>'


class ParserTests(unittest.TestCase):
    def test_global_deferred_and_aliases(self):
        # Promise 2 resolves at global index 3, whose key/text refs are 4/5.
        wire = '[{"_1":2},"data",["P",2]]\nP2:[{"_4":5},"text","中文\\n😀"]'
        self.assertEqual(hydrate(wire), {'data': {'text': '中文\n😀'}})
        self.assertEqual(hydrate('[["P",0]]\nP0:[{"_2":3},"text","a"]'), {'text':'a'})
        self.assertEqual(hydrate('[[1,2,3],"first",["P",2],["P",3]]\nP2:["shared"]\nP3:[["Z",4]]'), ['first','shared','shared'])
        self.assertEqual(hydrate('[[-5,-4,-7,-1]]'), [None,-0.0,None,None])

    def test_bad_streams_fail_explicitly(self):
        for wire in ['[]','{}','[[1]]','[["P",0]]','[["P",0]]\nE0:["denied"]',
                     '[[999]]','[["X",0]]','[[-2]]', '[["P",0]]\nP0:[]',
                     '[["P",0]]\nP0:["one"]\nP0:["two"]', '[{}]\nBAD']:
            with self.subTest(wire=wire), self.assertRaises(ExtractionError): hydrate(wire)

    def test_classification_preserves_every_record(self):
        msgs = [message('u'), message('a','assistant'), message('s','system'), message('t','tool'),
                message('h','user',metadata={'is_visually_hidden_from_conversation':True}),
                message('i','assistant',channel='analysis'),
                message('c','assistant',recipient='container.exec'),message('e',text='')]
        result=normalize(fixture(msgs))
        self.assertEqual([m['id'] for m in result['messages']],['u','a'])
        self.assertEqual(len(result['records']),8)
        self.assertEqual(result['report']['raw_messages'],8)
        self.assertEqual(sum(result['report']['categories'].values()),8)
        self.assertEqual(result['messages'][0]['create_time'],1234567890.5)
        self.assertEqual(result['messages'][0]['text_sha256'],hashlib.sha256(TEXT.encode()).hexdigest())
        self.assertFalse(result['report']['verified_against_live_page'])
        self.assertEqual(result['report']['reference_comparison'],'not_performed')

    def test_non_text_duplicate_and_unfinished_not_silent(self):
        msgs=[message('a'),message('a','assistant',status='in_progress'),
              message('media',content={'content_type':'multimodal_text','parts':['caption',{'type':'image','asset_pointer':'synthetic'}]})]
        result=normalize(fixture(msgs))
        self.assertEqual(result['report']['status'],'partial')
        self.assertEqual(len(result['report']['issues']),3)
        self.assertEqual(result['messages'][-1]['text'],'caption')
        self.assertEqual(result['records'][-1]['message']['content'],msgs[-1]['content'])

    def test_current_branch_and_linear_disagreement(self):
        nodes={'r':{'id':'r','parent':None,'message':message('u')},
               'a':{'id':'a','parent':'r','message':message('a','assistant')},
               'b':{'id':'b','parent':'r','message':message('b','assistant')}}
        conv={'mapping':nodes,'current_node':'a'}
        result=normalize(conv)
        self.assertEqual(result['report']['off_branch_nodes'],1)
        self.assertEqual([m['id'] for m in result['messages']],['u','a'])
        for bad in [dict(conv,current_node='missing'),dict(conv,linear_conversation=[nodes['b']])]:
            with self.assertRaises(ExtractionError):normalize(bad)
        nodes['r']['parent']='a'
        with self.assertRaises(ExtractionError):normalize(conv)

    def test_multichunk_exact_text(self):
        result=parse_html(html_for())
        self.assertEqual([m['text'] for m in result['messages']],[TEXT,TEXT])
        json.dumps(result,allow_nan=False)
        with self.assertRaises(ExtractionError): parse_html('streamController.enqueue("unterminated')


class FakeResponse:
    def __init__(self,data,declared=None,content_type='text/html'):
        from email.message import Message
        self.body=BytesIO(data);self.headers=Message();self.headers['Content-Type']=content_type
        if declared is not None:self.headers['Content-Length']=str(declared)
    def read1(self,size): return self.body.read(size)
    def __enter__(self):return self
    def __exit__(self,*args):pass


class NetworkTests(unittest.TestCase):
    def opener(self,response):
        mock=patch('chatkeep.urllib.request.build_opener').start()
        self.addCleanup(patch.stopall)
        mock.return_value.open.return_value=response
        return mock

    def test_url_security(self):
        for url in ['http://chatgpt.com/share/'+URL.rsplit('/',1)[1],URL+'?x=1',URL+'#x',URL+'/',
                    URL.replace('chatgpt.com','chatgpt.com.evil.test'),URL.replace('chatgpt.com','user@chatgpt.com'),
                    URL.replace('chatgpt.com','chatgpt.com:443'),URL.replace('11111111-','1111111--',1),
                    'https://127.0.0.1/share/'+URL.rsplit('/',1)[1]]:
            with self.subTest(url=url),self.assertRaises(ExtractionError):chatkeep.validate_url(url)

    def test_download_size_content_type_and_truncation(self):
        for response,max_bytes in [(FakeResponse(b'abcdef'),5),(FakeResponse(b'a',declared=6),5),
                                    (FakeResponse(b'a',declared=3),10),(FakeResponse(b'{}',content_type='application/json'),10)]:
            self.opener(response)
            with self.assertRaises(ExtractionError):fetch_html(URL,max_bytes=max_bytes)
            patch.stopall()
        self.opener(FakeResponse(b'<html>ok</html>'))
        self.assertEqual(fetch_html(URL),'<html>ok</html>')

    def test_timeout_http_errors_and_redirects(self):
        for error in [socket.timeout(),urllib.error.HTTPError(URL,403,'denied',{},None),urllib.error.URLError('offline')]:
            mock=self.opener(FakeResponse(b''));mock.return_value.open.side_effect=error
            with self.assertRaises(ExtractionError):fetch_html(URL)
            patch.stopall()
        # Exercise the real redirect guard through the opener handler list.
        actual=urllib.request.build_opener
        captured=[]
        def capture(handler):captured.append(handler);return actual(handler)
        with patch('chatkeep.urllib.request.build_opener',side_effect=capture), patch('urllib.request.OpenerDirector.open',return_value=FakeResponse(b'ok')):
            fetch_html(URL)
        for target in ['https://evil.test/', 'http://chatgpt.com/', 'https://chatgpt.com/api/private', URL+'?x=1']:
            with self.assertRaises(ExtractionError):captured[0]().redirect_request(urllib.request.Request(URL),None,302,'',{},target)


class WordTests(unittest.TestCase):
    def test_reopen_and_exact_text_order(self):
        model=parse_html(html_for())
        content=build_docx(model, rich=False)
        doc=Document(BytesIO(content))
        bodies=[p.text for i,p in enumerate(doc.paragraphs) if i and doc.paragraphs[i-1].style.name=='Heading 2']
        self.assertEqual(bodies,[TEXT,TEXT])
        self.assertEqual(doc.paragraphs[0].text,model['title'])
        self.assertTrue(content.startswith(b'PK'))

    def test_large_document_no_truncation(self):
        # 1,200 messages; >1M Unicode characters, alternating roles.
        texts=[('%04d 中文😀\n<box>完整正文</box>\n'%i)*50 for i in range(1200)]
        model=normalize(fixture([message(str(i),'user' if i%2==0 else 'assistant',t) for i,t in enumerate(texts)]))
        doc=Document(BytesIO(build_docx(model, rich=False)))
        bodies=[p.text for i,p in enumerate(doc.paragraphs) if i and doc.paragraphs[i-1].style.name=='Heading 2']
        self.assertEqual(bodies,texts)

    def test_filename_and_invalid_xml(self):
        self.assertEqual(filename_for('../CON/\\?'), '_CON___' + '.docx')
        self.assertEqual(filename_for('CON'),'_CON.docx')
        model=normalize(fixture([message('a',text='bad\x00text')]))
        with self.assertRaises(ExportError):build_docx(model, rich=False)


class HTTPTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server=Server(('127.0.0.1',0),Handler)
        cls.thread=threading.Thread(target=cls.server.serve_forever,daemon=True);cls.thread.start()
        cls.base='http://127.0.0.1:%d'%cls.server.server_port
    @classmethod
    def tearDownClass(cls):cls.server.shutdown();cls.server.server_close();cls.thread.join()
    def post(self,payload,headers=None):
        return urllib.request.urlopen(urllib.request.Request(self.base+'/api/convert',data=json.dumps(payload).encode(),headers=headers or {'Content-Type':'application/json'}),timeout=10)
    def test_complete_http_pipeline(self):
        with patch('chatkeep.fetch_html',return_value=html_for()), self.post({'url':URL}) as response:
            events=[json.loads(line) for line in response if line.strip()]
            self.assertEqual(events[0]['stage'], 'fetch')
            self.assertEqual(events[-1]['stage'], 'done')
            self.assertIn('parse', [e['stage'] for e in events])
            self.assertIn('word', [e['stage'] for e in events])
            self.assertIn('pack', [e['stage'] for e in events])
            for e in events:
                if e.get('percent') is not None:
                    self.assertGreaterEqual(e['percent'], 0)
                    self.assertLessEqual(e['percent'], 100)
            import base64
            doc=Document(BytesIO(base64.b64decode(events[-1]['document'])))
            self.assertIn('正文不能刪',document_text(doc))
            self.assertEqual(response.headers['Cache-Control'],'no-store')
    def test_error_not_success(self):
        with patch('chatkeep.fetch_html',side_effect=ExtractionError('擷取逾時')),self.post({'url':URL}) as response:
            events=[json.loads(line) for line in response if line.strip()]
            self.assertEqual(events[-1]['stage'],'error');self.assertNotIn('document',events[-1])
        for payload,headers,code in [({'url':'https://evil.test'},{'Content-Type':'application/json'},400),
                                     ({'url':URL},{'Content-Type':'application/json','Origin':'https://evil.test'},403),
                                     ({'url':URL},{'Content-Type':'text/plain'},403)]:
            with self.assertRaises(urllib.error.HTTPError) as caught:self.post(payload,headers)
            self.assertEqual(caught.exception.code,code)
    def test_static_and_unknown_download(self):
        for path in ['/','/app.js','/style.css']:
            with urllib.request.urlopen(self.base+path) as response:self.assertEqual(response.status,200)
        with self.assertRaises(urllib.error.HTTPError) as caught:urllib.request.urlopen(self.base+'/download/expired')
        self.assertEqual(caught.exception.code,404)
