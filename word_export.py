"""One editable Word builder, with actual OOXML admission at message boundaries."""
from io import BytesIO
import re
from lxml import etree
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from rich_word import render_message
from word_layout import configure, polish_tables
from export_limits import (CAPACITY, OBSERVER, MAX_MESSAGES, MAX_TEXT, MAX_MESSAGE, MAX_ARTIFACT)


class ExportError(ValueError):
    def __init__(self, message, code='EXPORT_LIMIT'):
        super().__init__(message)
        self.code = code


class VolumeFull(ExportError):
    def __init__(self):
        super().__init__('單卷已達安全 Word 結構容量。', 'VOLUME_CAPACITY')


class LimitedBuffer(BytesIO):
    def write(self, data):
        if self.tell() + len(data) > MAX_ARTIFACT:
            raise ExportError('輸出超過 16 MB 安全下載上限；未提供部分文件。', 'OUTPUT_LIMIT')
        return super().write(data)


def filename_for(title):
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f\x7f]', '_', str(title)).strip(' .')[:100]
    name = name.encode('utf-8')[:200].decode('utf-8', errors='ignore').rstrip(' .')
    if not name:
        name = '未命名對話'
    if re.fullmatch(r'(CON|PRN|AUX|NUL|COM[0-9¹²³]|LPT[0-9¹²³])', name.split('.')[0].rstrip(' '), re.I):
        name = '_' + name
    return name + '.docx'


def _check_xml(text):
    if any(not (c in '\t\n\r' or '\x20' <= c <= '\ud7ff' or '\ue000' <= c <= '\ufffd' or '\U00010000' <= c <= '\U0010ffff') for c in text):
        raise ExportError('來源含 Word 無法保存的控制字元；未產生可能缺字的文件。', 'INVALID_CHARACTERS')


def validate_model(model):
    if model.get('schema_version') != '1.0' or not model.get('messages'):
        raise ExportError('沒有可匯出的標準化訊息。')
    if len(str(model['title'])) > 1000:
        raise ExportError('對話標題超過安全長度。', 'INPUT_LIMIT')
    messages = model['messages']
    if len(messages) > MAX_MESSAGES or sum(len(m['text']) for m in messages) > MAX_TEXT:
        raise ExportError('對話超過總輸入安全上限（5,000 則或 500 萬字元），不截斷內容。', 'INPUT_LIMIT')
    if any(len(m['text']) > MAX_MESSAGE for m in messages):
        raise ExportError('單則訊息超過 100 萬字元；不能在訊息內切割。', 'MESSAGE_TOO_LARGE')
    ids = [m['id'] for m in messages if m.get('id')]
    if len(ids) != len(set(ids)):
        raise ExportError('有效輸出序列含重複訊息 ID，未輸出重複內容。', 'DUPLICATE_MESSAGE')


def measure(elements):
    nodes = size = cells = 0
    for element in elements:
        xml = etree.tostring(element, encoding='utf-8')
        # Child serialization repeats inherited namespaces; the DOCX root only
        # writes them once. Count the actual emitted subtree bytes without repeats.
        opening = xml[:xml.find(b'>')+1]
        size += len(xml) - sum(len(m.group()) for m in re.finditer(rb' xmlns(?::[\w]+)?="[^"]*"', opening))
        for node in element.iter():
            nodes += 1
            if node.tag == qn('w:tc'): cells += 1
    return nodes, size, cells


class WordVolume:
    def __init__(self, model, rich=True, capacity=CAPACITY, check=None, message_offset=0):
        self.model, self.rich, self.capacity, self.check = model, rich, capacity, check
        self.message_offset = message_offset
        self.document = configure(Document())
        title = str(model['title']); _check_xml(title)
        self.document.add_heading(title, 0)
        self.document.add_paragraph('ChatKeep｜對話收藏')
        self.document.add_paragraph('完整性：尚未完成逐訊息比對。可匯出訊息 %d 則。' % len(model['messages']))
        for warning in model['report']['warnings']:
            _check_xml(warning); self.document.add_paragraph(warning)
        initial = measure(list(self.document.element.body))
        self.stats = (initial[0]+16, initial[1]+8192, initial[2])  # reserve final volume label
        self.indices, self.warnings = [], []

    def _over(self, stats):
        return any(a > b for a, b in zip(stats, (self.capacity.elements, self.capacity.xml_bytes, self.capacity.cells)))

    def add(self, index):
        if self.check: self.check()
        body = self.document.element.body
        # New blocks are inserted before the final section properties. Keep one
        # stable sibling anchor instead of rescanning all prior messages (O(n²)).
        anchor = body[-2] if len(body)>1 else None
        old_rels = set(self.document.part.rels)
        pending_warnings = []
        calls = 0
        def added():
            result = []
            element = anchor.getnext() if anchor is not None else body[0]
            while element is not None and element.tag != qn('w:sectPr'):
                result.append(element)
                element = element.getnext()
            return result
        def observer(elements=0, cells=0):
            nonlocal calls
            calls += 1
            if self.check: self.check()
            if calls % 32 and elements < 1024 and cells < 100: return
            current = measure(added())
            combined = tuple(a+b for a,b in zip(self.stats, current))
            if self._over((combined[0]+elements, combined[1], combined[2]+cells)):
                raise VolumeFull()
        token = OBSERVER.set(observer)
        try:
            message = self.model['messages'][index]
            role = {'user': '使用者', 'assistant': 'ChatGPT'}.get(message['role'], message['role'])
            position = index + self.message_offset + 1
            paragraph = self.document.add_heading('%d · %s' % (position, role), 2)
            start = OxmlElement('w:bookmarkStart'); start.set(qn('w:id'), str(position)); start.set(qn('w:name'), 'ChatKeepMessage_%d' % position)
            end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(position))
            paragraph._p.append(start); paragraph._p.append(end)
            text = message['text']; _check_xml(text)
            if self.rich:
                for warning in render_message(self.document, text):
                    observer()
                    pending_warnings.append({'message_index': index, 'reason': warning})
                    self.document.add_paragraph('［轉換警告］'+warning, style='ChatKeep Warning')
            else:
                observer(elements=2+2*(text.count('\n')+text.count('\r')+text.count('\t')))
                self.document.add_paragraph().add_run(text)
            if message.get('category') == 'non_text':
                self.document.add_paragraph('［此訊息含未還原的非文字內容；以上保留可取得文字。］')
            # Polish just the tables added by this message, including nested cell endings.
            polish_tables(self.document, tables=[Table(e, self.document._body) for e in added() if e.tag == qn('w:tbl')])
            current = measure(added())
            combined = tuple(a+b for a,b in zip(self.stats, current))
            if self._over(combined): raise VolumeFull()
            self.stats = combined
            self.indices.append(index+self.message_offset)
            self.warnings.extend(pending_warnings)
        except BaseException:
            for element in added(): body.remove(element)
            for rel in set(self.document.part.rels)-old_rels:
                del self.document.part.rels[rel]
            raise
        finally:
            OBSERVER.reset(token)

    def pack(self, number=1, multipart=False, on_pack=None):
        if self.check: self.check()
        if multipart:
            label = self.model.get('volume_label_prefix','') + '第%d卷' % number
            self.document.paragraphs[0].text = '%s｜%s' % (self.model['title'], label)
            self.document.add_paragraph('%s｜有效輸出訊息位置 %d–%d，共 %d 則。' % (label, self.indices[0]+1, self.indices[-1]+1, len(self.indices)))
        if on_pack: on_pack()
        data = LimitedBuffer()
        self.document.save(data)
        if self.check: self.check()
        self.stats = measure(list(self.document.element.body))
        if self._over(self.stats): raise VolumeFull()
        return data.getvalue()


def build_docx(model, rich=True, on_progress=None, on_pack=None, warnings_out=None):
    """Backward compatible single-DOCX API; the artifact core handles rollover."""
    validate_model(model)
    volume = WordVolume(model, rich=rich)
    total = len(model['messages'])
    if on_progress: on_progress(0, total)
    for index in range(total):
        try: volume.add(index)
        except VolumeFull as exc:
            if not volume.indices:
                raise ExportError('單則訊息超過安全 Word 結構容量，不能從訊息中切割。', 'MESSAGE_TOO_LARGE') from exc
            raise ExportError('對話需要分卷，請使用支援 DOCX／ZIP 的新版 API。', 'MULTIPART_REQUIRED') from exc
        if on_progress and ((index+1) % 10 == 0 or index+1 == total): on_progress(index+1, total)
    if warnings_out is not None: warnings_out.extend(volume.warnings)
    return volume.pack(on_pack=on_pack)
