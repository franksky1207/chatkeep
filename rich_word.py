"""Editable Word rendering for safe, static ChatGPT UI and Markdown.

No screenshots; all text is native Word runs or table cell content.
"""
import re
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from semantic import Node, parse, flatten

HEAD = re.compile(r"^(#{1,6})\s+(.+)$")
BULLET = re.compile(r"^(\s*)([-+*]|\d+[.)])\s+(.+)$")
BUBBLE = re.compile(r"^\s*(?:(LINE)\s*[｜|:]\s*)?([^：:\n]{1,20})[：:]\s*(.+)$")
CLOCK = re.compile(r"^\d{1,2}:\d{2}(?:\s|$)")
DIVIDER = re.compile(r"^\s*[-*_]{3,}\s*$")
CELL_LIMIT = 12

def shade(cell, fill="EDF4EF"):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)

def celltext(cell, value):
    cell.text = value
    for p in cell.paragraphs:
        p.paragraph_format.widow_control = True
        p.paragraph_format.space_after = Pt(1)

def box(document, text, fill="EFF4F2"):
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade(cell, fill)
    celltext(cell, text)
    return table

def bubble(document, speaker, content, timestamp=""):
    table = document.add_table(rows=1, cols=2)
    mine = speaker in ("我", "玩家", "使用者", "You")
    active = table.cell(0, 1 if mine else 0)
    shade(active, "E1F3DF" if mine else "F1F3F5")
    celltext(active, (speaker + (" · " + timestamp if timestamp else "") + "\n" + content))
    table.cell(0, 0 if mine else 1).text = ""

def markdown(document, value):
    lines = value.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if DIVIDER.match(line):
            document.add_paragraph("──────────")
            i += 1
            continue
        heading = HEAD.match(line)
        if heading:
            document.add_heading(heading.group(2), level=min(len(heading.group(1))+1, 6))
            i += 1
            continue
        if line.startswith(">"):
            box(document, line.lstrip("> "), "F3F3F3")
            i += 1
            continue
        if line.count("|") >= 2 and i+1 < len(lines) and re.fullmatch(r"[\s|:\-]+", lines[i+1] or "") and "-" in lines[i+1]:
            rows = [line]
            i += 2
            while i < len(lines) and "|" in lines[i]:
                rows.append(lines[i])
                i += 1
            matrix = [[cell.strip() for cell in row.strip().strip("|").split("|")] for row in rows]
            width = max(map(len, matrix))
            if width <= CELL_LIMIT:
                table = document.add_table(rows=len(matrix), cols=width)
                table.style = "Table Grid"
                for r, cols in enumerate(matrix):
                    for c, val in enumerate(cols):
                        celltext(table.cell(r,c), val)
                        if r == 0: shade(table.cell(r,c))
            else:
                for row in rows: document.add_paragraph(row)
            continue
        bullet = BULLET.match(line)
        if bullet:
            numbered = bullet.group(2)[0].isdigit()
            try:
                document.add_paragraph(bullet.group(3), style="List Number" if numbered else "List Bullet")
            except KeyError:
                document.add_paragraph(line)
            i += 1
            continue
        match = BUBBLE.match(line)
        if match and (match.group(1) or CLOCK.match(line) or "LINE" in value[:80]):
            bubble(document, match.group(2).strip(), match.group(3))
            i += 1
            continue
        # Keep consecutive prose lines in one editable paragraph, preserving line breaks.
        group = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not (HEAD.match(lines[i]) or BULLET.match(lines[i]) or DIVIDER.match(lines[i]) or lines[i].startswith(">") or (BUBBLE.match(lines[i]) and "LINE" in value[:80])):
            group.append(lines[i])
            i += 1
        document.add_paragraph("\n".join(group))

def render(document, node):
    kind = node.kind
    if kind in ("markdown", "text"):
        markdown(document, flatten(node) if node.children else node.text)
    elif kind == "code":
        p = document.add_paragraph()
        p.style = document.styles["Normal"]
        p.paragraph_format.left_indent = Pt(12)
        p.paragraph_format.right_indent = Pt(8)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.widow_control = True
        run = p.add_run(node.text)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(38, 59, 70)
    elif kind == "asset":
        attrs = node.attrs
        label = attrs.get("alt") or attrs.get("query") or "圖片原始資產未包含在分享資料"
        box(document, "[圖片未嵌入] " + label, "F5F4EC")
    elif kind == "decor":
        # Decorative icons are not independently meaningful.
        return
    elif kind in ("box", "card"):
        # Extract all descendant text in source order. Nested UI syntax removed,
        # but raw text is not discarded.
        box(document, flatten(node))
    elif kind == "row":
        children = [c for c in node.children if flatten(c).strip()]
        if len(children) >= 2 and len(children) <= 4:
            table = document.add_table(rows=1, cols=len(children))
            table.style = "Table Grid"
            for i, child in enumerate(children):
                celltext(table.cell(0,i), flatten(child))
        else:
            box(document, flatten(node))
    elif kind in ("col", "caption", "text", "bold", "italic"):
        value = flatten(node)
        p = document.add_paragraph()
        r = p.add_run(value)
        if kind=="caption": r.font.size=Pt(9); r.font.color.rgb=RGBColor(100,110,110)
        if kind=="bold": r.bold=True
        if kind=="italic": r.italic=True
    else:
        markdown(document, flatten(node))

def render_message(document, value):
    nodes, warnings = parse(value)
    for node in nodes:
        render(document, node)
    return warnings
