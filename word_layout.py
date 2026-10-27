"""Reusable, editable Word reading layout. No image rasterization or page-per-turn."""
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

FONT = "Microsoft JhengHei"

def _font(style, size, color=None, bold=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    if color:
        style.font.color.rgb = RGBColor(*color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), FONT)

def _keep_rows(table):
    for row in table.rows:
        trpr = row._tr.get_or_add_trPr()
        # Long cards must flow across pages instead of becoming an oversized,
        # unsplittable row. Short bubbles/rows still stay together.
        for old in list(trpr.findall(qn('w:cantSplit'))):
            trpr.remove(old)
        size = sum(len(cell.text) for cell in row.cells)
        lines = sum(cell.text.count('\n') for cell in row.cells)
        if size <= 800 and lines <= 12 and not any(cell.tables for cell in row.cells):
            trpr.append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.widow_control = True

def _field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    command = OxmlElement("w:instrText")
    command.set(qn("xml:space"), "preserve")
    command.text = " " + instruction + " "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, command, separate, placeholder, end):
        run._r.append(element)

def configure(document):
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.45)
    section.right_margin = Cm(2.45)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    normal = document.styles["Normal"]
    _font(normal, 10.5, (40, 50, 57))
    fmt = normal.paragraph_format
    fmt.space_after = Pt(5)
    fmt.line_spacing = 1.28
    fmt.widow_control = True

    title = document.styles["Title"]
    _font(title, 20, (29, 66, 60), True)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.keep_with_next = True

    specs = [
        ("Heading 1", 15, (24, 67, 63)),
        ("Heading 2", 12, (35, 83, 77)),
        ("Heading 3", 11, (49, 83, 79)),
        ("Heading 4", 10.5, (60, 86, 83)),
    ]
    for name, size, color in specs:
        style = document.styles[name]
        _font(style, size, color, True)
        style.paragraph_format.space_before = Pt(11 if name != "Heading 2" else 8)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
    for name in ("Heading 5", "Heading 6"):
        document.styles[name].paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        document.styles[name].paragraph_format.space_after = Pt(2)

    warning = document.styles.add_style('ChatKeep Warning', WD_STYLE_TYPE.PARAGRAPH)
    _font(warning, 9, (110, 80, 40))
    warning.paragraph_format.space_after = Pt(4)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footrun = footer.add_run("ChatKeep  ·  ")
    footrun.font.size = Pt(8)
    footrun.font.color.rgb = RGBColor(118, 130, 130)
    _field(footer, "PAGE")
    return document

def polish_tables(document):
    for table in document.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        _keep_rows(table)
        for row in table.rows:
            for cell in row.cells:
                polish_tables(cell)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(1)
                    paragraph.paragraph_format.line_spacing = 1.2
